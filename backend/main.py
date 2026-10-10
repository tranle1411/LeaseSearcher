import io
import logging
import os
import time
import uuid
from contextlib import asynccontextmanager
from datetime import date
from pathlib import Path

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.middleware.gzip import GZipMiddleware
from openpyxl import Workbook

from backend.schemas import Search
from backend.observability import LimitBody, configure, request_id as request_context
from backend.scoring import recommend
from backend.storage import engine_for, read_snapshot, select_dataset

log = logging.getLogger("leasesearcher")
CUISINES = (
    "Vietnamese",
    "Chinese",
    "Japanese",
    "Korean",
    "Thai",
    "Indian",
    "Filipino",
    "Indonesian",
    "Malaysian",
    "Taiwanese",
    "Other",
)


@asynccontextmanager
async def lifespan(app):
    configure(log)
    app.state.dataset = read_snapshot()
    app.state.engine = engine_for(os.environ["DATABASE_URL"]) if os.getenv("DATABASE_URL") else None
    app.state.last_read = 0.0
    app.state.cached_read = (app.state.dataset, "snapshot")
    try:
        yield
    finally:
        if app.state.engine is not None:
            app.state.engine.dispose()


app = FastAPI(title="LeaseSearcher", version="2.0.0", lifespan=lifespan)
app.add_middleware(LimitBody)
app.add_middleware(GZipMiddleware, minimum_size=1000, compresslevel=5)


@app.middleware("http")
async def observe(request: Request, call_next):
    request_id = uuid.uuid4().hex
    started = time.perf_counter()
    try:
        length = int(request.headers.get("content-length", "0"))
    except ValueError:
        return JSONResponse({"detail": "Invalid content length"}, status_code=400)
    if length > 16384:
        return JSONResponse({"detail": "Request too large"}, status_code=413)
    token = request_context.set(request_id)
    try:
        response = await call_next(request)
    except Exception as error:
        log.error("request_failed", extra={"error_type": type(error).__name__})
        response = JSONResponse(
            {"detail": "Service error. Please retry.", "request_id": request_id}, status_code=500
        )
    response.headers["X-Request-ID"] = request_id
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    if request.url.path.startswith("/assets/") and response.status_code == 200:
        response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    elif request.url.path == "/":
        response.headers["Cache-Control"] = "no-cache"
    log.info(
        "request_complete",
        extra={
            "method": request.method,
            "path": request.url.path,
            "status": response.status_code,
            "duration_ms": round((time.perf_counter() - started) * 1000, 2),
        },
    )
    request_context.reset(token)
    return response


def current():
    now = time.monotonic()
    if app.state.engine is not None and now - app.state.last_read < 30:
        return app.state.cached_read
    try:
        selected = select_dataset(app.state.dataset, app.state.engine)
    except Exception as error:
        log.warning("database_snapshot_fallback", extra={"error_type": type(error).__name__})
        selected = app.state.dataset, "snapshot-fallback"
    if app.state.engine is not None:
        app.state.cached_read = selected
        app.state.last_read = now
    return selected


@app.get("/api/health")
def health():
    return {"status": "ok", "version": app.state.dataset.version}


@app.get("/api/metadata")
def metadata():
    dataset, mode = current()
    return {
        "version": dataset.version,
        "collection_period": dataset.collection_period,
        "counties": sorted({row.county for row in dataset.listings}),
        "cuisines": CUISINES,
        "listing_count": len(dataset.listings),
        "mode": mode,
        "report": dataset.report,
        "sources": dataset.sources,
        "demographics_available": bool(dataset.areas),
        "eligibility_policy": dataset.eligibility_policy,
        "pricing_policy": dataset.pricing_policy,
        "allow_zip_centroids": dataset.allow_zip_centroids,
    }


def results(search):
    dataset, mode = current()
    if search.dataset_version and search.dataset_version != dataset.version:
        raise HTTPException(409, "Dataset changed. Run your search again before exporting.")
    if search.evaluated_on and search.evaluated_on != date.today():
        raise HTTPException(
            409, "Verification date changed. Run your search again before exporting."
        )
    if search.cuisine not in CUISINES:
        raise HTTPException(422, "Choose a supported cuisine")
    if search.county and search.county not in {row.county for row in dataset.listings}:
        raise HTTPException(422, "Choose a county in this dataset")
    result = recommend(dataset, search, mode, CUISINES)
    log.info(
        "search_complete",
        extra={
            "dataset_version": dataset.version,
            "scoring_version": result.scoring_version,
            "eligible": result.total_eligible,
            "exclusions": result.exclusions,
            "mode": mode,
        },
    )
    return result


@app.post("/api/search")
def search(body: Search):
    return results(body)


def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@app.post("/api/export")
def export(body: Search):
    result = results(body)
    book = Workbook()
    sheet = book.active
    sheet.title = "Top recommendations"
    sheet.append(
        [
            "Rank",
            "Address",
            "County",
            "Fit score",
            "Monthly cost used (USD)",
            "Site",
            "Location",
            "Competition",
            "Evidence coverage %",
            "Listing URL",
            "Listing ID",
            "Verified through",
            "Eligibility evidence",
            "Warnings",
            "Cost basis",
            "Estimated size (sq ft)",
            "Coordinate precision",
            "Advertised price",
        ]
    )
    for match in result.matches:
        sheet.append(
            [
                safe_cell(value)
                for value in [
                    match.rank,
                    match.listing.address,
                    match.listing.county,
                    match.fit_score,
                    match.effective_monthly_cost,
                    *match.components.values(),
                    match.evidence_coverage,
                    str(match.listing.source_url),
                    match.listing.id,
                    match.listing.verification.valid_until.isoformat()
                    if match.listing.verification
                    else "Unverified",
                    match.listing.verification.evidence
                    if match.listing.verification
                    else "Demo: existence and restaurant eligibility assumed",
                    " | ".join(match.warnings),
                    match.cost_basis,
                    match.assumed_sqft,
                    match.listing.coordinate_precision,
                    match.listing.advertised_price,
                ]
            ]
        )
        sheet.cell(sheet.max_row, 10).hyperlink = str(match.listing.source_url)
        sheet.cell(sheet.max_row, 10).style = "Hyperlink"
    method = book.create_sheet("Method and exclusions")
    method.append(["Dataset", result.dataset_version])
    method.append(["Scoring", result.scoring_version])
    method.append(["Evaluated on", result.evaluated_on.isoformat()])
    method.append(["Eligible", result.total_eligible])
    method.append(["Eligibility policy", result.eligibility_policy])
    method.append(["Pricing policy", result.pricing_policy])
    method.append(
        [
            "Demo estimates",
            "Advertised rate and compatible size use range midpoints. Annual rates are divided by 12. Disclosed monthly fees are added; undisclosed fees are unknown and excluded. Exact unit size and restaurant eligibility are assumed, not verified.",
        ]
    )
    method.append(
        [
            "Caution",
            "Historical listings; no assurance of current availability or business success.",
        ]
    )
    method.append(["Exclusions", result.exclusion_rule])
    for key, value in result.exclusions.items():
        method.append([key, value])
    for key, value in body.model_dump(mode="json").items():
        method.append([key, safe_cell(str(value))])
    for tab in book:
        tab.freeze_panes = "A2"
        tab.auto_filter.ref = tab.dimensions
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = min(
                70, max(16, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    stream = io.BytesIO()
    book.save(stream)
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="lease-shortlist.xlsx"'},
    )


dist = Path("frontend/dist")
if dist.exists():
    app.mount("/assets", StaticFiles(directory=dist / "assets"), name="assets")

    @app.get("/")
    def index():
        return FileResponse(dist / "index.html")
