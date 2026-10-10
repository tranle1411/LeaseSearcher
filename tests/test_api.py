from io import BytesIO

from openpyxl import load_workbook


def test_api_health_and_search(client):
    assert client.get("/api/health").json()["status"] == "ok"
    response = client.post("/api/search", json={})
    assert response.status_code == 200
    assert response.headers["x-request-id"]
    body = response.json()
    assert len(body["matches"]) == 1
    assert body["matches"][0]["evidence_coverage"] == 100
    assert client.post("/api/search", json={"cuisine": "not-supported"}).status_code == 422
    assert client.post("/api/search", json={"county": "not-a-county"}).status_code == 422
    assert client.post("/api/search", json={"budget": -1}).status_code == 422
    assert client.post("/api/search", json={"unknown": 1}).status_code == 422
    assert client.post("/api/search", json={"dataset_version": "old"}).status_code == 409


def test_export_matches_search_and_guards_version(client):
    result = client.post("/api/search", json={}).json()
    response = client.post("/api/export", json={"dataset_version": result["dataset_version"]})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.active.cell(2, 2).value == result["matches"][0]["listing"]["address"]
    assert book.active.cell(2, 4).value == result["matches"][0]["fit_score"]
    assert client.post("/api/export", json={"dataset_version": "old"}).status_code == 409
    assert client.post("/api/export", json={"evaluated_on": "2000-01-01"}).status_code == 409
    assert book.active.cell(2, 11).value == result["matches"][0]["listing"]["id"]
    assert book.active.cell(2, 10).hyperlink.target == result["matches"][0]["listing"]["source_url"]
    zero = client.post("/api/export", json={"budget": 1})
    assert load_workbook(BytesIO(zero.content)).active.max_row == 1


def test_database_outage_uses_same_version(client, monkeypatch):
    from backend import main

    version = client.get("/api/health").json()["version"]

    def broken(*args):
        raise ConnectionError("test outage")

    monkeypatch.setattr(main, "select_dataset", broken)
    response = client.post("/api/search", json={}).json()
    assert response["dataset_version"] == version
    assert response["mode"] == "snapshot-fallback"


def test_excel_formula_is_text_and_request_size_limited(client, dataset):
    from backend.main import app

    listing = dataset.listings[0].model_copy(update={"address": '=HYPERLINK("bad")'})
    app.state.dataset = dataset.model_copy(update={"listings": (listing,)})
    response = client.post("/api/export", json={})
    cell = load_workbook(BytesIO(response.content)).active.cell(2, 2)
    assert cell.data_type == "s"
    assert client.post("/api/search", content="x" * 17000).status_code == 413
    assert client.post("/api/search", content=iter([b"x" * 9000, b"x" * 9000])).status_code == 413


def test_database_retry_cooldown(client, monkeypatch, dataset):
    from backend import main

    calls = []

    class UnavailableEngine:
        def dispose(self):
            pass

    main.app.state.engine = UnavailableEngine()

    def broken(*args):
        calls.append(1)
        raise ConnectionError("test outage")

    monkeypatch.setattr(main, "select_dataset", broken)
    assert client.post("/api/search", json={}).json()["mode"] == "snapshot-fallback"
    assert client.post("/api/search", json={}).json()["mode"] == "snapshot-fallback"
    assert len(calls) == 1
    main.app.state.last_read = 0
    monkeypatch.setattr(main, "select_dataset", lambda *args: (dataset, "postgresql"))
    assert client.post("/api/search", json={}).json()["mode"] == "postgresql"
