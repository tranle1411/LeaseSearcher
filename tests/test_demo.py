from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.schemas import Search
from backend.scoring import exclusion_reasons, quoted_cost, recommend


def demo_listing(listing, **updates):
    return listing.model_copy(
        update={
            "verification": None,
            "rent_basis": "per_sqft_month",
            "rent_low": 2,
            "rent_high": 4,
            **updates,
        }
    )


def test_midpoint_cost_unknown_fees_and_explicit_rejection(dataset, listing):
    row = demo_listing(listing)
    assert quoted_cost(row, Search(), "demo_assumed", "midpoint") == (
        3750,
        1250,
        "advertised_estimate",
    )
    assert quoted_cost(row, Search(), "demo_assumed", "conservative")[0] == 6000
    data = dataset.model_copy(
        update={
            "listings": (row,),
            "eligibility_policy": "demo_assumed",
            "pricing_policy": "midpoint",
        }
    )
    result = recommend(data, Search(), "test", ["Vietnamese"])
    assert result.matches[0].listing.verification is None
    assert any("fees are unknown" in s for s in result.matches[0].assumptions)
    assert (
        quoted_cost(
            demo_listing(listing, disclosed_monthly_fees=200), Search(), "demo_assumed", "midpoint"
        )[0]
        == 3950
    )
    denied = row.model_copy(
        update={"verification": listing.verification.model_copy(update={"eligible": False})}
    )
    assert exclusion_reasons(denied, Search(), date.today(), "demo_assumed", "midpoint")


@pytest.mark.parametrize(
    "updates",
    [
        {"lat": None},
        {"sqft_min": None},
        {"rent_basis": None},
        {"source_url": None},
        {"availability": "unavailable"},
    ],
)
def test_demo_missing_required_details_still_excluded(listing, updates):
    assert exclusion_reasons(
        demo_listing(listing, **updates), Search(), date.today(), "demo_assumed", "midpoint"
    )


def test_estimate_respects_budget_size_and_coordinate_policy(listing):
    row = demo_listing(listing, coordinate_precision="zip_centroid")
    assert exclusion_reasons(row, Search(), date.today(), "demo_assumed", "midpoint")
    assert not exclusion_reasons(row, Search(), date.today(), "demo_assumed", "midpoint", True)
    assert "Above monthly budget" in exclusion_reasons(
        row, Search(budget=3749), date.today(), "demo_assumed", "midpoint", True
    )
    assert quoted_cost(row, Search(sqft_min=2000), "demo_assumed", "midpoint")[0] is None
    # Estimate is the midpoint of the compatible offered size range.
    assert quoted_cost(row, Search(sqft_max=1200), "demo_assumed", "midpoint")[1] == 1100
    flat = demo_listing(listing, rent_basis="monthly", rent_low=2000, rent_high=4000)
    assert quoted_cost(flat, Search(), "demo_assumed", "midpoint")[0] == 3000


def test_demo_export_without_verification(client, dataset, listing):
    from backend.main import app

    app.state.dataset = dataset.model_copy(
        update={
            "listings": (demo_listing(listing),),
            "eligibility_policy": "demo_assumed",
            "pricing_policy": "midpoint",
        }
    )
    result = client.post("/api/search", json={}).json()
    assert result["matches"][0]["effective_monthly_cost"] == 3750
    response = client.post("/api/export", json={"dataset_version": result["dataset_version"]})
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.cell(2, 5).value == 3750
    assert sheet.cell(2, 12).value == "Unverified"
    assert sheet.cell(2, 15).value == "advertised_estimate"
