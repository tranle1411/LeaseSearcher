from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from pipeline.review import FIELDS, read_reviews, template


def test_review_requires_explicit_evidence_and_booleans(dataset, tmp_path):
    snapshot = tmp_path / "snapshot.json"
    snapshot.write_text(dataset.model_dump_json(), encoding="utf-8")
    workbook = tmp_path / "review.xlsx"
    template(snapshot, workbook)
    assert read_reviews(workbook) == {}
    book = load_workbook(workbook)
    row = {
        "eligible": True,
        "complete_fees": True,
        "lease_space_confirmed": True,
        "monthly_occupancy_cost": 3000,
        "verified_on": date.today().isoformat(),
        "valid_until": (date.today() + timedelta(days=10)).isoformat(),
        "verified_by": "Test reviewer",
        "evidence": "TEST ONLY source reference",
    }
    for key, value in row.items():
        book["Review"].cell(2, FIELDS.index(key) + 1, value)
    book.save(workbook)
    assert read_reviews(workbook)[dataset.listings[0].id]["verification"]["eligible"]
    book["Review"].cell(2, FIELDS.index("complete_fees") + 1, "probably")
    book.save(workbook)
    with pytest.raises(ValueError, match="explicitly"):
        read_reviews(workbook)
