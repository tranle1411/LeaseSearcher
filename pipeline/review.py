"""Export a consistent broker-review workbook; import only explicitly completed reviews."""

import argparse
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from backend.schemas import Dataset, Verification

FIELDS = [
    "listing_id",
    "address_reference",
    "source_url_reference",
    "advertised_rent_reference",
    "eligible",
    "complete_fees",
    "lease_space_confirmed",
    "monthly_occupancy_cost",
    "verified_on",
    "valid_until",
    "verified_by",
    "evidence",
    "unit",
    "sqft_min",
    "sqft_max",
    "lat",
    "lon",
    "coordinate_source",
    "parking",
    "availability",
    "availability_verified_on",
]


def template(snapshot, output):
    dataset = Dataset.model_validate_json(Path(snapshot).read_text(encoding="utf-8"))
    book = Workbook()
    sheet = book.active
    sheet.title = "Review"
    sheet.append(FIELDS)
    for listing in dataset.listings:
        address = listing.address
        if address.startswith(("=", "+", "-", "@")):
            address = "'" + address
        sheet.append(
            [listing.id, address, str(listing.source_url or ""), listing.monthly_base_rent]
        )
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    sheet.column_dimensions["B"].width = 60
    sheet.column_dimensions["L"].width = 65
    notes = book.create_sheet("Instructions")
    for text in [
        "Only complete rows with actual evidence. Blank evidence rows are ignored.",
        "eligible = restaurant zoning/permit eligibility checked for this unit and proposed use.",
        "complete_fees = all recurring occupancy costs are known, including rent and applicable fees.",
        "lease_space_confirmed = exact offered unit, area and position checked; no building averages.",
        "Use TRUE or FALSE for booleans. verified_on/valid_until use YYYY-MM-DD.",
        "evidence must identify the document/source and what was verified. verified_by identifies reviewer.",
        "Enter reviewed size and precise coordinates if legacy values are missing or inaccurate.",
        "availability: unconfirmed, confirmed or unavailable; confirmed requires a date.",
        "Do not enter passwords, private contact details or confidential documents in a public snapshot.",
        f"Source dataset: {dataset.version}; advertised rent is reference only, never a verified total.",
    ]:
        notes.append([text])
    notes.column_dimensions["A"].width = 125
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(f"Review template: {output}")


def boolean(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().lower() in ("true", "false"):
        return value.strip().lower() == "true"
    raise ValueError("Boolean review fields must explicitly be TRUE or FALSE")


def date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def read_reviews(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = book["Review"].iter_rows(values_only=True)
        headers = next(rows)
        if list(headers) != FIELDS:
            raise ValueError("Use the generated review template without changing columns")
        overrides = {}
        for number, values in enumerate(rows, 2):
            row = dict(zip(headers, values))
            if not row.get("evidence"):
                continue
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                raise ValueError(
                    f"Row {number}: review fields must be literal values, not formulas"
                )
            identity = row["listing_id"]
            if identity in overrides:
                raise ValueError(f"Row {number}: duplicate reviewed listing ID")
            try:
                verification = Verification(
                    eligible=boolean(row["eligible"]),
                    complete_fees=boolean(row["complete_fees"]),
                    lease_space_confirmed=boolean(row["lease_space_confirmed"]),
                    monthly_occupancy_cost=row["monthly_occupancy_cost"],
                    verified_on=date_value(row["verified_on"]),
                    valid_until=date_value(row["valid_until"]),
                    verified_by=row["verified_by"],
                    evidence=row["evidence"],
                )
                override = {"verification": verification.model_dump(mode="json")}
                for key in FIELDS[12:]:
                    if row[key] is not None:
                        override[key] = (
                            boolean(row[key]) if key == "parking" else date_value(row[key])
                        )
                overrides[identity] = override
            except ValueError as error:
                raise ValueError(f"Review row {number}: {error}") from error
        return overrides
    finally:
        book.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snapshot", default="datasets/snapshot.json")
    parser.add_argument("--output", default=".local/broker-review.xlsx")
    args = parser.parse_args()
    template(args.snapshot, args.output)
